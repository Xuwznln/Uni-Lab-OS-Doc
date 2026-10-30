"""OS 本地模式的 Backend 同形试剂（Reagent）HTTP 适配器。"""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from pathlib import PurePath
from typing import Any, Callable, Dict, List, Optional, Sequence, Type
from uuid import UUID
from zipfile import BadZipFile

from fastapi import APIRouter, File, Query, UploadFile
from fastapi.responses import JSONResponse
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from unilabos.app.scheduler.inventory.backend_contract import (
    INVALID_PARAMETER,
    BackendContractError,
)
from unilabos.app.scheduler.inventory.backend_response import call, success
from unilabos.app.scheduler.inventory.reagent_contract import BackendReagentService

_MAX_IMPORT_ROWS = 500
_MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024
_OPTIONAL_CSV_FIELDS = {
    "name_en",
    "aliases",
    "molecular_formula",
    "smiles",
    "inchi_key",
    "molecular_weight",
    "density_g_per_ml",
    "description",
    "meta_data",
    "reagent_info_uuid",
    "concentration_value",
    "concentration_unit",
    "source",
    "observed_at",
    "container_capacity",
    "expected_material_revision",
}


class ReagentModel(BaseModel):
    """忽略 Backend DTO 的未知扩展字段，保留向前兼容。"""

    model_config = ConfigDict(extra="ignore")


class ReagentInfoCreateRequest(ReagentModel):
    cas: str = ""
    name: str
    name_en: Optional[str] = None
    aliases: List[str] = Field(default_factory=list)
    molecular_formula: Optional[str] = None
    smiles: Optional[str] = None
    inchi_key: Optional[str] = None
    molecular_weight: Optional[float] = None
    density_g_per_ml: Optional[float] = None
    physical_state: str = "unknown"
    description: Optional[str] = None
    meta_data: Dict[str, Any] = Field(default_factory=dict)


class ReagentInfoUpdateRequest(ReagentModel):
    cas: Optional[str] = None
    name: Optional[str] = None
    name_en: Optional[str] = None
    aliases: Optional[List[str]] = None
    molecular_formula: Optional[str] = None
    smiles: Optional[str] = None
    inchi_key: Optional[str] = None
    molecular_weight: Optional[float] = None
    density_g_per_ml: Optional[float] = None
    physical_state: Optional[str] = None
    description: Optional[str] = None
    meta_data: Optional[Dict[str, Any]] = None


class MaterialReagentRequest(ReagentModel):
    """创建容器时可内联携带的试剂字段，不包含容器物料身份。"""

    reagent_info_uuid: Optional[UUID] = None
    cas: str = ""
    physical_state: str = "unknown"
    density_g_per_ml: Optional[float] = None
    concentration_value: Optional[float] = None
    concentration_unit: Optional[str] = None
    quantity: float
    quantity_unit: str
    source: Optional[str] = None
    observed_at: Optional[datetime] = None
    description: Optional[str] = None
    meta_data: Dict[str, Any] = Field(default_factory=dict)
    container_capacity: Optional[Dict[str, Any]] = None
    expected_material_revision: Optional[int] = None


class ReagentCreateRequest(MaterialReagentRequest):
    """为已存在容器单独创建试剂实例的请求。"""

    material_uuid: UUID


class ReagentInfoBatchRequest(ReagentModel):
    """一次导入多条试剂身份；默认整批原子提交。"""

    # 先保留原始行，再由统一导入器逐行校验；这样 atomic=false 才能返回
    # 某一行错误而不让 Pydantic 在整个请求边界提前拒绝整批。
    items: List[Dict[str, Any]] = Field(
        min_length=1, max_length=_MAX_IMPORT_ROWS
    )
    atomic: bool = True


class ReagentBatchRequest(ReagentModel):
    """一次导入多条容器级试剂；默认整批原子提交。"""

    items: List[Dict[str, Any]] = Field(
        min_length=1, max_length=_MAX_IMPORT_ROWS
    )
    atomic: bool = True


class ReagentUpdateRequest(ReagentModel):
    concentration_value: Optional[float] = None
    concentration_unit: Optional[str] = None
    quantity: float
    quantity_unit: str
    source: Optional[str] = None
    observed_at: Optional[datetime] = None
    expected_revision: Optional[int] = None
    description: Optional[str] = None
    meta_data: Dict[str, Any] = Field(default_factory=dict)
    container_capacity: Optional[Dict[str, Any]] = None
    expected_material_revision: Optional[int] = None


def _import_error(
    message: str,
    *,
    details: Dict[str, Any],
    code: int = INVALID_PARAMETER,
) -> JSONResponse:
    """构造带行号明细的批量导入错误信封。"""

    return JSONResponse(
        status_code=200,
        content={
            "code": code,
            "error": {"msg": message, "details": details},
        },
    )


def _parse_json_value(
    value: str, *, field: str, row_number: int, source: str = "CSV"
) -> Any:
    """解析表格中的 JSON 字段，并把错误定位到具体行列。"""

    try:
        return json.loads(value)
    except json.JSONDecodeError as error:
        raise BackendContractError(
            INVALID_PARAMETER,
            f"{source} 第 {row_number} 行的 {field} 不是有效 JSON",
        ) from error


def _normalize_tabular_row(
    row: Dict[str, Any], *, row_number: int, allowed_fields: set[str], source: str = "CSV"
) -> Dict[str, Any]:
    """把表格单元格转换为 JSON DTO 可接受的值。"""

    unknown = sorted(set(row) - allowed_fields)
    if unknown:
        raise BackendContractError(
            INVALID_PARAMETER,
            f"{source} 第 {row_number} 行包含未知字段：{', '.join(unknown)}",
        )
    normalized: Dict[str, Any] = {}
    for key, raw in row.items():
        if raw is None:
            if key == "aliases":
                normalized[key] = []
            elif key == "meta_data":
                normalized[key] = {}
            elif key == "cas":
                normalized[key] = ""
            else:
                normalized[key] = None
            continue
        value = str(raw).strip()
        if value == "":
            if key == "aliases":
                normalized[key] = []
            elif key == "meta_data":
                normalized[key] = {}
            elif key == "cas":
                normalized[key] = ""
            else:
                normalized[key] = None if key in _OPTIONAL_CSV_FIELDS else ""
        elif key in {"aliases", "meta_data", "container_capacity"}:
            parsed = _parse_json_value(
                value, field=key, row_number=row_number, source=source
            )
            if key == "aliases" and not isinstance(parsed, list):
                raise BackendContractError(
                    INVALID_PARAMETER,
                    f"{source} 第 {row_number} 行的 aliases 必须是 JSON 数组",
                )
            if key in {"meta_data", "container_capacity"} and not isinstance(parsed, dict):
                raise BackendContractError(
                    INVALID_PARAMETER,
                    f"{source} 第 {row_number} 行的 {key} 必须是 JSON 对象",
                )
            normalized[key] = parsed
        else:
            normalized[key] = value
    return normalized


def _is_blank_cell(value: Any) -> bool:
    """判断 Excel 单元格是否为空；数值 0 不能被当作空值。"""

    return value is None or (isinstance(value, str) and not value.strip())


def _parse_xlsx_file(
    content: bytes,
    *,
    allowed_fields: set[str],
) -> List[Dict[str, Any]]:
    """读取 XLSX 的第一个工作表，并转换成统一的导入行。"""

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as error:  # pragma: no cover - 安装包声明保证正常环境具备依赖。
        raise BackendContractError(
            INVALID_PARAMETER,
            "XLSX 导入需要安装 openpyxl 依赖",
        ) from error

    try:
        workbook = load_workbook(
            filename=io.BytesIO(content), read_only=True, data_only=True
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as error:
        raise BackendContractError(INVALID_PARAMETER, "XLSX 导入文件格式无效") from error
    try:
        worksheet = workbook.worksheets[0] if workbook.worksheets else None
        if worksheet is None:
            raise BackendContractError(INVALID_PARAMETER, "XLSX 文件没有工作表")
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration as error:
            raise BackendContractError(INVALID_PARAMETER, "XLSX 文件没有表头") from error

        headers = [str(value).strip() if value is not None else "" for value in header_values]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not field for field in headers):
            raise BackendContractError(INVALID_PARAMETER, "XLSX 文件必须包含非空表头")
        if len(set(headers)) != len(headers):
            raise BackendContractError(INVALID_PARAMETER, "XLSX 表头不能重复")
        unknown = sorted(set(headers) - allowed_fields)
        if unknown:
            raise BackendContractError(
                INVALID_PARAMETER,
                f"XLSX 表头包含未知字段：{', '.join(unknown)}",
            )

        parsed_rows: List[Dict[str, Any]] = []
        for row_number, values in enumerate(rows, 2):
            if not any(not _is_blank_cell(value) for value in values):
                continue
            if len(values) > len(headers) and any(
                not _is_blank_cell(value) for value in values[len(headers) :]
            ):
                raise BackendContractError(
                    INVALID_PARAMETER,
                    f"XLSX 第 {row_number} 行列数多于表头",
                )
            row = {
                headers[index]: values[index] if index < len(values) else None
                for index in range(len(headers))
            }
            parsed_rows.append(
                _normalize_tabular_row(
                    row,
                    row_number=row_number,
                    allowed_fields=allowed_fields,
                    source="XLSX",
                )
            )
            if len(parsed_rows) > _MAX_IMPORT_ROWS:
                raise BackendContractError(
                    INVALID_PARAMETER,
                    f"单次最多导入 {_MAX_IMPORT_ROWS} 行",
                )
        if not parsed_rows:
            raise BackendContractError(INVALID_PARAMETER, "导入文件没有数据行")
        return parsed_rows
    finally:
        workbook.close()


def _parse_import_file(
    content: bytes,
    *,
    filename: str,
    content_type: str | None,
    allowed_fields: set[str],
) -> List[Dict[str, Any]]:
    """解析 JSON/CSV 导入文件为字典行，并限制大小和行数。"""

    suffix = PurePath(filename.lower()).suffix
    media_type = (content_type or "").split(";", 1)[0].strip().lower()
    is_xlsx = suffix == ".xlsx" or media_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    if is_xlsx:
        return _parse_xlsx_file(content, allowed_fields=allowed_fields)

    is_json = suffix == ".json" or media_type in {
        "application/json",
        "application/ld+json",
    }
    is_csv = suffix in {".csv", ".tsv"} or media_type in {
        "text/csv",
        "text/tab-separated-values",
        "application/vnd.ms-excel",
    }
    if not is_json and not is_csv:
        raise BackendContractError(
            INVALID_PARAMETER,
            "仅支持 .json、.csv、.tsv 或 .xlsx 文件",
        )
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise BackendContractError(INVALID_PARAMETER, "导入文件必须使用 UTF-8 编码") from error

    if is_json:
        try:
            payload = json.loads(text)
        except json.JSONDecodeError as error:
            raise BackendContractError(INVALID_PARAMETER, "JSON 导入文件格式无效") from error
        rows = payload.get("items") if isinstance(payload, dict) else payload
        if not isinstance(rows, list):
            raise BackendContractError(
                INVALID_PARAMETER,
                "JSON 导入文件必须是数组，或包含 items 数组的对象",
            )
        if any(not isinstance(row, dict) for row in rows):
            raise BackendContractError(INVALID_PARAMETER, "JSON 导入数组的每一项必须是对象")
        if not rows:
            raise BackendContractError(INVALID_PARAMETER, "导入文件没有数据行")
        if len(rows) > _MAX_IMPORT_ROWS:
            raise BackendContractError(
                INVALID_PARAMETER,
                f"单次最多导入 {_MAX_IMPORT_ROWS} 行",
            )
        return [dict(row) for row in rows]

    delimiter = "\t" if suffix == ".tsv" else ","
    reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
    fieldnames = [str(field or "").strip() for field in (reader.fieldnames or [])]
    if not fieldnames or any(not field for field in fieldnames):
        raise BackendContractError(INVALID_PARAMETER, "CSV 文件必须包含非空表头")
    if len(set(fieldnames)) != len(fieldnames):
        raise BackendContractError(INVALID_PARAMETER, "CSV 表头不能重复")
    unknown = sorted(set(fieldnames) - allowed_fields)
    if unknown:
        raise BackendContractError(
            INVALID_PARAMETER,
            f"CSV 表头包含未知字段：{', '.join(unknown)}",
        )
    rows: List[Dict[str, Any]] = []
    for row_number, row in enumerate(reader, 2):
        if None in row:
            raise BackendContractError(
                INVALID_PARAMETER,
                f"CSV 第 {row_number} 行列数多于表头",
            )
        if not any(str(value or "").strip() for value in row.values()):
            continue
        rows.append(
            _normalize_tabular_row(
                {str(key).strip(): value for key, value in row.items()},
                row_number=row_number,
                allowed_fields=allowed_fields,
            )
        )
        if len(rows) > _MAX_IMPORT_ROWS:
            raise BackendContractError(
                INVALID_PARAMETER,
                f"单次最多导入 {_MAX_IMPORT_ROWS} 行",
            )
    if not rows:
        raise BackendContractError(INVALID_PARAMETER, "导入文件没有数据行")
    return rows


def _validation_errors(error: ValidationError) -> List[Dict[str, str]]:
    """将 Pydantic 错误压缩成稳定的字段/消息结构。"""

    return [
        {
            "field": ".".join(str(part) for part in detail.get("loc", ())) or "row",
            "message": str(detail.get("msg") or "字段无效"),
        }
        for detail in error.errors()
    ]


def _validated_rows(
    rows: Sequence[Dict[str, Any]],
    model: Type[BaseModel],
) -> tuple[List[tuple[int, Dict[str, Any]]], List[Dict[str, Any]]]:
    """逐行执行 DTO 校验，返回可写行与行级错误。"""

    valid: List[tuple[int, Dict[str, Any]]] = []
    errors: List[Dict[str, Any]] = []
    for row_number, row in enumerate(rows, 1):
        try:
            parsed = model.model_validate(row)
        except ValidationError as error:
            errors.append(
                {"row": row_number, "errors": _validation_errors(error)}
            )
            continue
        valid.append((row_number, parsed.model_dump(mode="json")))
    return valid, errors


def _batch_create(
    service: BackendReagentService,
    rows: Sequence[Dict[str, Any]],
    *,
    model: Type[BaseModel],
    batch_create: Callable[[List[Dict[str, Any]]], List[Dict[str, Any]]],
    single_create: Callable[[Dict[str, Any]], Dict[str, Any]],
    atomic: bool,
) -> JSONResponse:
    """统一执行 JSON/文件批量导入并返回行级结果。"""

    valid, errors = _validated_rows(rows, model)
    total = len(rows)
    if atomic and errors:
        return _import_error(
            "批量导入校验失败，未写入任何数据",
            details={"total": total, "created": 0, "failed": len(errors), "errors": errors},
        )
    if not valid:
        return _import_error(
            "批量导入没有可写入的数据",
            details={"total": total, "created": 0, "failed": len(errors), "errors": errors},
        )

    created: List[Dict[str, Any]] = []
    if atomic:
        try:
            created = batch_create([payload for _, payload in valid])
        except BackendContractError as error:
            batch_error = {
                "row": None,
                "errors": [{"message": error.message}],
            }
            return _import_error(
                error.message,
                details={
                    "total": total,
                    "created": 0,
                    "failed": total,
                    "errors": [*errors, batch_error],
                },
                code=error.code,
            )
    else:
        for row_number, payload in valid:
            try:
                created.append(single_create(payload))
            except BackendContractError as error:
                errors.append({"row": row_number, "errors": [{"message": error.message}]})
    result = {
        "total": total,
        "created": len(created),
        "failed": len(errors),
        "atomic": atomic,
        "items": created,
        "errors": errors,
    }
    return success(result)


async def _read_import_rows(
    upload: UploadFile,
    *,
    allowed_fields: set[str],
) -> List[Dict[str, Any]]:
    """读取上传文件并交给统一解析器。"""

    content = await upload.read(_MAX_IMPORT_FILE_BYTES + 1)
    if len(content) > _MAX_IMPORT_FILE_BYTES:
        raise BackendContractError(INVALID_PARAMETER, "导入文件不能超过 5 MiB")
    return _parse_import_file(
        content,
        filename=upload.filename or "",
        content_type=upload.content_type,
        allowed_fields=allowed_fields,
    )


def create_reagent_router(service: BackendReagentService) -> APIRouter:
    """创建化学身份、试剂实例、CAS 查询和台账路由。

    参数：``service`` 是绑定当前 ``inventory.db`` 的领域服务。返回：挂在
    ``/api/v1`` 父路由下的无前缀 Router。异常：装配错误原样抛出。
    """

    router = APIRouter(tags=["backend-reagent-contract"])

    @router.get("/compounds/{cas}")
    def lookup_compound(cas: str) -> JSONResponse:
        return call(service.lookup_compound, cas)

    @router.post("/reagent-infos")
    def create_reagent_info(body: ReagentInfoCreateRequest) -> JSONResponse:
        return call(
            service.create_reagent_info,
            body.model_dump(mode="json"),
            status_code=201,
        )

    @router.get("/reagent-infos")
    def list_reagent_infos(
        page: int = Query(default=0),
        page_size: int = Query(default=0),
        name: str = Query(default=""),
        cas: str = Query(default=""),
        physical_state: str = Query(default=""),
    ) -> JSONResponse:
        return call(
            service.list_reagent_infos,
            page=page,
            page_size=page_size,
            name=name,
            cas=cas,
            physical_state=physical_state,
        )

    reagent_info_fields = set(ReagentInfoCreateRequest.model_fields)

    @router.post("/reagent-infos/batch")
    def create_reagent_infos_batch(body: ReagentInfoBatchRequest) -> JSONResponse:
        """批量登记试剂身份；默认整批原子提交。"""

        return _batch_create(
            service,
            body.items,
            model=ReagentInfoCreateRequest,
            batch_create=service.create_reagent_infos,
            single_create=service.create_reagent_info,
            atomic=body.atomic,
        )

    @router.post("/reagent-infos/import")
    async def import_reagent_infos(
        file: UploadFile = File(...),
        atomic: bool = Query(default=True),
    ) -> JSONResponse:
        """从 JSON/CSV/TSV 文件批量登记试剂身份。"""

        try:
            rows = await _read_import_rows(file, allowed_fields=reagent_info_fields)
        except BackendContractError as error:
            return _import_error(
                error.message,
                details={"total": 0, "created": 0, "failed": 0, "errors": []},
            )
        return _batch_create(
            service,
            rows,
            model=ReagentInfoCreateRequest,
            batch_create=service.create_reagent_infos,
            single_create=service.create_reagent_info,
            atomic=atomic,
        )

    @router.get("/reagent-infos/{reagent_info_uuid}")
    def get_reagent_info(reagent_info_uuid: UUID) -> JSONResponse:
        return call(service.get_reagent_info, str(reagent_info_uuid))

    @router.get("/reagent-infos/{reagent_info_uuid}/structure-3d")
    def get_reagent_info_structure(reagent_info_uuid: UUID) -> JSONResponse:
        return call(service.get_reagent_info_structure, str(reagent_info_uuid))

    @router.put("/reagent-infos/{reagent_info_uuid}")
    def update_reagent_info(
        reagent_info_uuid: UUID, body: ReagentInfoUpdateRequest
    ) -> JSONResponse:
        return call(
            service.update_reagent_info,
            str(reagent_info_uuid),
            body.model_dump(mode="json", exclude_unset=True),
        )

    @router.delete("/reagent-infos/{reagent_info_uuid}")
    def delete_reagent_info(reagent_info_uuid: UUID) -> JSONResponse:
        return call(service.delete_reagent_info, str(reagent_info_uuid))

    @router.post("/reagents")
    def create_reagent(body: ReagentCreateRequest) -> JSONResponse:
        return call(
            service.create_reagent,
            body.model_dump(mode="json"),
            status_code=201,
        )

    reagent_fields = set(ReagentCreateRequest.model_fields)

    @router.post("/reagents/batch")
    def create_reagents_batch(body: ReagentBatchRequest) -> JSONResponse:
        """批量登记容器级试剂；默认整批原子提交。"""

        return _batch_create(
            service,
            body.items,
            model=ReagentCreateRequest,
            batch_create=service.create_reagents,
            single_create=service.create_reagent,
            atomic=body.atomic,
        )

    @router.post("/reagents/import")
    async def import_reagents(
        file: UploadFile = File(...),
        atomic: bool = Query(default=True),
    ) -> JSONResponse:
        """从 JSON/CSV/TSV 文件批量登记容器级试剂。"""

        try:
            rows = await _read_import_rows(file, allowed_fields=reagent_fields)
        except BackendContractError as error:
            return _import_error(
                error.message,
                details={"total": 0, "created": 0, "failed": 0, "errors": []},
            )
        return _batch_create(
            service,
            rows,
            model=ReagentCreateRequest,
            batch_create=service.create_reagents,
            single_create=service.create_reagent,
            atomic=atomic,
        )

    @router.get("/reagents")
    def list_reagents(
        page: int = Query(default=0),
        page_size: int = Query(default=0),
        material_uuid: Optional[UUID] = Query(default=None),
        reagent_info_uuid: Optional[UUID] = Query(default=None),
        keyword: str = Query(default=""),
        cas: str = Query(default=""),
        barcode: str = Query(default=""),
    ) -> JSONResponse:
        return call(
            service.list_reagents,
            page=page,
            page_size=page_size,
            material_uuid=str(material_uuid) if material_uuid else "",
            reagent_info_uuid=(
                str(reagent_info_uuid) if reagent_info_uuid else ""
            ),
            keyword=keyword,
            cas=cas,
            barcode=barcode,
        )

    @router.get("/reagents/{reagent_uuid}")
    def get_reagent(reagent_uuid: UUID) -> JSONResponse:
        return call(service.get_reagent, str(reagent_uuid))

    @router.put("/reagents/{reagent_uuid}")
    def update_reagent(
        reagent_uuid: UUID, body: ReagentUpdateRequest
    ) -> JSONResponse:
        return call(
            service.update_reagent,
            str(reagent_uuid),
            body.model_dump(mode="json", exclude_unset=True),
        )

    @router.delete("/reagents/{reagent_uuid}")
    def delete_reagent(reagent_uuid: UUID) -> JSONResponse:
        return call(service.delete_reagent, str(reagent_uuid))

    @router.get("/reagent-history/{history_uuid}")
    def get_reagent_history(history_uuid: UUID) -> JSONResponse:
        return call(service.get_reagent_history, str(history_uuid))

    @router.get("/materials/{material_uuid}/reagent-history")
    def list_reagent_history(
        material_uuid: UUID,
        page: int = Query(default=0),
        page_size: int = Query(default=0),
    ) -> JSONResponse:
        return call(
            service.list_reagent_history,
            str(material_uuid),
            page=page,
            page_size=page_size,
        )

    return router


__all__ = [
    "MaterialReagentRequest",
    "ReagentBatchRequest",
    "ReagentInfoBatchRequest",
    "create_reagent_router",
]
